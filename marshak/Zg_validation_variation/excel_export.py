import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


def df_to_excel_with_gradient(df, output_path, pct_columns=None):
    if pct_columns is None:
        pct_columns = []
        for col in df.columns:
            if pd.api.types.is_numeric_dtype(df[col]):
                if df[col].min() >= 0 and df[col].max() <= 1:
                    pct_columns.append(col)

    df_export = df.copy()

    for col in df_export.columns:
        if pd.api.types.is_string_dtype(df_export[col]):
            df_export[col] = df_export[col].str.replace("_", " ")

    # df_export["lc_class"] = df_export["lc_class"].str.replace("_", " ")

    for col in pct_columns:
        if col in df_export.columns:
            df_export[col] = (df_export[col] * 100).round(2)

    df_export.to_excel(output_path, index=False)

    wb = load_workbook(output_path)
    ws = wb.active

    for col in pct_columns:
        if col not in df.columns:
            continue

        col_idx = df.columns.get_loc(col) + 1
        col_letter = get_column_letter(col_idx)

        for row_idx in range(2, len(df) + 2):
            cell = ws[f"{col_letter}{row_idx}"]
            value = cell.value

            if value is not None and isinstance(value, (int, float)):
                intensity = max(0, min(100, value)) / 100
                r = int(255 * (1 - intensity))
                g = int(255 * (1 - intensity))
                b = 255

                hex_color = f"{r:02X}{g:02X}{b:02X}"
                cell.fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

    wb.save(output_path)
