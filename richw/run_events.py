#! /usr/bin/env -S python -i

### /u/aurora-r0/richw/pkgs/miniforge3/envs/dist-s1-research/bin/python -i

# Setup for and conditionally run Dist-S1 results and pull Dist-HLS results

import sys
import os
import inspect
from pathlib import Path
import contextily as ctx
from datetime import datetime, timedelta

import math
from shapely.geometry import Point
from shapely.geometry import box
from geopy.distance import geodesic
import pandas as pd
import geopandas as gpd
from osgeo import gdal
import rasterio
from rasterio.windows import Window
from rasterio.transform import from_bounds
from rasterio.crs import CRS
from rasterio.warp import calculate_default_transform, reproject, Resampling
from dem_stitcher.rio_tools import reproject_arr_to_match_profile
from PIL import Image
import numpy as np
import warnings
import copy
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches
import leafmap
import cv2
import asf_search

import dist_s1
from dist_s1.packaging import (
  convert_geotiff_to_png
)   
from dist_s1.dist_processing import (
    merge_burst_metrics_and_serialize
)
from dist_s1.workflows import(
    run_dist_s1_workflow,
    run_dist_s1_sas_prep_workflow
)
from dist_s1.constants import (
    BASE_DATE_FOR_CONFIRMATION,
    DISTLABEL2VAL,
    DIST_STATUS_CMAP,
    TIF_LAYERS,
    TIF_LAYER_DTYPES,
    TIF_LAYER_NODATA_VALUES,
)        
from dist_s1_enumerator import (
    enumerate_one_dist_s1_product,
    get_burst_ids_in_mgrs_tiles,
    get_burst_table_from_mgrs_tiles,
    get_lut_by_mgrs_tile_ids,
    get_mgrs_tiles_overlapping_geometry,
    get_rtc_s1_ts_metadata_from_mgrs_tiles,
    get_mgrs_table,
)
from dist_s1_enumerator.mgrs_burst_data import (
    get_mgrs_tile_table_by_ids,
    get_burst_ids_in_mgrs_tiles
)
from dist_s1_enumerator.dist_enum import enumerate_dist_s1_products

from dist_s1.rio_tools import open_one_ds
from distmetrics.rio_tools import merge_with_weighted_overlap
        
import plot_fcns
import utils_geotif 
import util_fcns
import debug_defs

from util_fcns import (
    show_mgrs,
    show_mgrs1,
    enumerate_post_dates,
    run_mgrs_seq_local,
    prod_from_dir,
    pull_dist_hls_seq_local,
    pull_hls_seq_local,
    pull_radd_seq_local,
)

def main() -> None:
  redo_radd_pull = False
  redo_dist_hls_pull = False
  redo_hls_pull = False
  redo_enumerate_s1 = False
  redo_run_s1 = False
  redo_s1_pngs = True
  redo_merge = True
  redo_dist_hls_pngs = True
  redo_hls_pngs = True
  redo_radd_pngs = False
  redo_pair_pngs = True
  redo_quad_pngs = True
  use_subset = True
  sas_prep_only = False
  dist_hls_prods = ['VEG-DIST-STATUS','GEN-DIST-STATUS']
  hls_prods = ['B02','B03','B04','jpg']
  radd_prods = ['jpg','tif']
  # circle marker radius as proportion of subset image width
  circ_size = 0.01
  geo_crs = CRS.from_epsg(4326)

  asf_search.constants.INTERNAL.CMR_TIMEOUT = 180

  # Directories
  basedir = '/home/richw/dat/opera/dist-s1/events'
  s1_dir = 's1'
  hls_dir = 'hls'
  dist_hls_dir = 'dist-hls'
  radd_dir = 'radd'

  # Ensemble of targets to analyze
  # Baltimore construction in 18SUJ is seen by Dist-HLS not much by Dist-S1
  target_names = ['tornado1','tornado2','logging1','mining1',
      'radd-road-const','radd-small-forestchg','radd-small-logging',
      'achafalaya-delta'] 
  event_dates = ['2025-05-16','2025-03-15','2025-01-01','2025-01-01',
      '2024-11-11','2024-01-01','2025-02-14','2026-02-21']
  start_dates = ['2025-04-16','2025-02-13','2024-01-01','2024-01-01',
      '2024-09-01','2024-01-01','2024-12-01','2025-06-21']
  stop_dates = ['2025-10-15','2025-08-14','2025-12-31','2025-12-31',
      '2025-12-31','2026-01-01','2025-09-01','2026-02-21']
  minlats = [36.93, 33.17, 46.28, -13.06, -1.456, 5.66, 7.560, 29.38]
  minlons = [-84.6, -88.05, -122.55, -70.17, -48.451, 39.21, 34.860, -91.55]
  maxlats = [37.14, 33.34, 46.52, -12.80, -1.415, 5.83, 7.585, 29.6]
  maxlons = [-84.0, -87.90, -122.16, -69.95, -48.343, 39.37, 34.910, -91.2]
  ptlats = [37.05, 33.27, 46.384, -13.006, -1.438, 5.749, 7.572, 29.46]
  ptlons = [-84.35, -87.96, -122.443, -70.059, -48.406, 39.29, 34.881, -91.44]
  sel_mgrs = [1,1,0,0,0,0,1,0]
  mp4_rates = [2,2,5,5,3,5,3,2]

  for itarget,target_name in enumerate(target_names):
    print(target_name)
    mp4_rate = mp4_rates[itarget]
    # Paths
    target_path = Path(basedir) / Path(target_name)
    target_pathstr = str(target_path)
    s1_pathbase = target_path / Path(s1_dir)
    dist_hls_pathbase = target_path / Path(dist_hls_dir)
    hls_pathbase = target_path / Path(hls_dir)
    radd_pathbase = target_path / Path(radd_dir)
    s1_pathbase.mkdir(parents=True, exist_ok=True)
    hls_pathbase.mkdir(parents=True, exist_ok=True)
    dist_hls_pathbase.mkdir(parents=True, exist_ok=True)
    radd_pathbase.mkdir(parents=True, exist_ok=True)

    # Target bounding box
    minlat = minlats[itarget]
    minlon = minlons[itarget]
    maxlat = maxlats[itarget]
    maxlon = maxlons[itarget]
    bbox,gdf_bbox = util_fcns.bbox_setup(minlat,minlon,maxlat,maxlon)
    gdf_mgrs_tiles = get_mgrs_tiles_overlapping_geometry(bbox)
    basename = target_pathstr + '/mgrs_' + target_name
    plot_fcns.plot_gdf_geoms(basename,
        [gdf_mgrs_tiles,gdf_bbox],['mgrs_tile_id','target box'])
        
    # Points of interest
    lat1 = ptlats[itarget]
    lon1 = ptlons[itarget]
    point1 = Point(lon1,lat1)
    gdf1 = gpd.GeoDataFrame({'Name': ['Point of interest'],
        'geometry': [point1]}, crs="EPSG:4326")

    # Select MGRS tile
    imgrs = sel_mgrs[itarget]
    mgrs_tile_id = gdf_mgrs_tiles['mgrs_tile_id'][imgrs]
    mask = gdf_mgrs_tiles['mgrs_tile_id'] == mgrs_tile_id
    gdf_mgrs1 = gdf_mgrs_tiles[mask]

    # Establish post dates
    event_date = event_dates[itarget]
    event_dt = datetime.strptime(event_date,"%Y-%m-%d")
    start_date = start_dates[itarget]
    stop_date = stop_dates[itarget]
    start_dt = datetime.strptime(start_date,"%Y-%m-%d")
    stop_dt = datetime.strptime(stop_date,"%Y-%m-%d")
    #start_dt = event_dt - timedelta(days=30)
    #stop_dt = event_dt + timedelta(days=60)
    #start_date = start_dt.strftime("%Y-%m-%d")
    #stop_date = stop_dt.strftime("%Y-%m-%d")

    # Select subset to run
    if itarget != 0:
        continue

    #raise Exception("stop for mgrs tile")

    # Determine post dates available to run
    filepath = s1_pathbase / 'df_sel_prod_post.parquet'
    if redo_enumerate_s1:
        sorted_post_dates,df_sel_prod_post = enumerate_post_dates(
            mgrs_tile_id,start_dt,stop_dt)
        df_sel_prod_post.to_parquet(filepath)
    else:
        df_sel_prod_post = gpd.read_parquet(filepath)
        post_dates = df_sel_prod_post.acq_date_for_mgrs_pass.unique()
        sorted_post_dates = np.sort(post_dates.astype('datetime64'))

    track_nums = df_sel_prod_post.track_number.unique()

    print(f"mgrs_tile_id: {mgrs_tile_id}")
    print(f"track_nums: {track_nums}")
    print(f"post dates: {sorted_post_dates}")
    print(f"sas_prep_only: {sas_prep_only}")

    name = target_pathstr + '/bbox_points'
    plot_fcns.plot_gdf_geoms(name,
        [gdf_mgrs1,gdf_bbox,gdf1],['mgrs_tile_id','target box','Point'])

    # Run the selected set of post dates
    if redo_run_s1:
        df_s1_prod = run_mgrs_seq_local(mgrs_tile_id,sorted_post_dates,
            df_sel_prod_post,sas_prep_only,str(s1_pathbase),bbox)
    df_s1_prod = prod_from_dir(str(s1_pathbase))

    # Pull DIST-HLS products for same date range
    if redo_dist_hls_pull:
        pull_dist_hls_seq_local(start_date,stop_date,
            str(dist_hls_pathbase),bbox,mgrs_tile_id,dist_hls_prods)
    df_dist_hls_prod = prod_from_dir(str(dist_hls_pathbase),
        prod_base="OPERA_L3_DIST-ALERT-HLS")

    # Pull HLS products for same date range
    if redo_hls_pull:
        pull_hls_seq_local(start_date,stop_date,
            str(hls_pathbase),bbox,mgrs_tile_id,hls_prods)
    df_hls_prod = prod_from_dir(str(hls_pathbase),
        prod_base="HLS.S30")

    # Pull RADD products for same date range
    if redo_radd_pull:
        pull_radd_seq_local(start_date,stop_date,
            str(radd_pathbase),bbox,mgrs_tile_id,radd_prods,'xx')
    df_radd_prod = prod_from_dir(str(radd_pathbase),
        prod_base="HLS.S30")

    prod_names_s1 = ['GEN-DIST-STATUS.tif',
        'GEN-METRIC.tif']
    prod_names_dist_hls = ['VEG-DIST-STATUS.tif']
    prod_names_hls = ['B03.tif']
    copol_names = ['VV','HH']
    crosspol_names = ['VH','HV']

    # Use first status product to set uniform resolution
    filename1 = Path(df_s1_prod.prod_name[0]) / (
        Path(df_s1_prod.prod_name[0]).name + '_' + prod_names_s1[0])
    dst_res = utils_geotif.projected_res_to_deg_res(filename1)

    root = s1_pathbase / Path(mgrs_tile_id + '_subset')
    burst_ids = [
        p.stem.split("_")[3] for p in root.rglob("*")
        if p.suffix == ".tif" and len(p.stem.split("_")) > 3
    ]
    unique_burst_ids = list(set(burst_ids))
            
    # Colormaps
    DIST_STATUS_CMAP = {
        0: (18, 18, 18, 255),  # No disturbance
        1: (0, 85, 85, 255),  # First low
        2: (137, 127, 78, 255),  # Provisional low
        3: (222, 224, 67, 255),  # Confrimed low
        4: (0, 136, 136, 255),  # First high
        5: (228, 135, 39, 255),  # Provisional high
        6: (224, 27, 7, 255),  # Confirmed high
        7: (119, 119, 119, 255),  # Confirmed low finished
        8: (221, 221, 221, 255),  # Confirmed high finished
        255: (0, 0, 0, 255),  # No data
    }

    status_pngs = []
    metric_pngs = []
    copol_pngs = []
    crosspol_pngs = []
    copol_values = []
    crosspol_values = []
    status = []
    metric = []
    dt_s1 = []
    for df_prod_row in df_s1_prod.itertuples(index=True):
        #if df_prod_row.Index == 0:
        print(df_prod_row.post_date)
        dt_s1.append(pd.to_datetime(df_prod_row.post_date))
        filename = Path(df_prod_row.prod_name) / (
            Path(df_prod_row.prod_name).name + '_' + prod_names_s1[0])
        with rasterio.open(filename) as src:
            value,row1,col1 = util_fcns.value_at(src,lon1,lat1)
            status.append(value)
        filename = Path(df_prod_row.prod_name) / (
            Path(df_prod_row.prod_name).name + '_' + prod_names_s1[1])
        with gdal.Open(filename) as ds:
            meta = ds.GetMetadata()
            opera_ids_row = meta['post_rtc_opera_ids'].split(",")
            burst_ids_row = [s.split("_")[3] for s in opera_ids_row]
            # But likely only one burst covering target point!
        with rasterio.open(filename) as src:
            value,row1,col1 = util_fcns.value_at(src,lon1,lat1)
            metric.append(value)

        ofile = Path(df_prod_row.prod_name) / 'dist_status.tif'
        ofile_png = ofile.with_suffix(".png")
        if redo_s1_pngs:
            # Reproject geotif into same destination geometry
            ifile = Path(df_prod_row.prod_name) / (
                Path(df_prod_row.prod_name).name + '_' + prod_names_s1[0])
            dst_width,dst_height = utils_geotif.reproject_to_geom(
                ifile,geo_crs,bbox,dst_res,ofile)
            # Make pngs
            titlestr = 'Dist-S1 ' + df_prod_row.post_date
            status_bounds2 = utils_geotif.geotif_to_png_map2(
                ofile,
                ofile_png,
                geo_crs,
                titlestr,
                DIST_STATUS_CMAP,
                0,6,lat1,lon1,circ_size,None)
        status_pngs.append(ofile_png)

        ofile = Path(df_prod_row.prod_name) / 'metric.tif'
        ofile_png = ofile.with_suffix(".png")
        if redo_s1_pngs:
            ifile = Path(df_prod_row.prod_name) / (
                Path(df_prod_row.prod_name).name + '_' + prod_names_s1[1])
            dst_width,dst_height = utils_geotif.reproject_to_geom(
                ifile,geo_crs,bbox,dst_res,ofile)
            # Make pngs
            titlestr = 'Dist-S1 Metric ' + df_prod_row.post_date
            metric_bounds2 = utils_geotif.geotif_to_png_map2(
                ofile,
                ofile_png,
                geo_crs,
                titlestr,
                'gray',
                0,6,lat1,lon1,circ_size,None)
        metric_pngs.append(ofile_png)

        if redo_s1_pngs:
            # Make mosaic of subset RTC data for post-dates
            parts2 = str(df_prod_row.prod_name).split("_")
            post_date_dt = datetime.strptime(parts2[4],"%Y%m%dT%H%M%SZ")
            post_date_str = post_date_dt.strftime("%Y-%m-%d")
            rtc_paths = [f for f in root.rglob("*")
                if post_date_str in str(f)]
            copol_paths = [p for p in rtc_paths
                if any(sub in p.name for sub in copol_names)]
            crosspol_paths = [p for p in rtc_paths
                if any(sub in p.name for sub in crosspol_names)]
            print('starting rtc_setup')
            titlestr = 'Copol ' + df_prod_row.post_date
            copol_tif,copol_png,copol_bounds = util_fcns.rtc_setup(
                copol_paths,mgrs_tile_id,bbox,dst_res,df_prod_row.prod_name,
                'copol',titlestr,0,0.3,lat1,lon1,circ_size,redo_merge)
            copol_pngs.append(copol_png)
            titlestr = 'Xpol ' + df_prod_row.post_date
            crosspol_tif,crosspol_png,crosspol_bounds = util_fcns.rtc_setup(
                crosspol_paths,mgrs_tile_id,bbox,dst_res,df_prod_row.prod_name,
                'crosspol',titlestr,0,0.1,lat1,lon1,circ_size,redo_merge)
            crosspol_pngs.append(crosspol_png)
            print('starting value extract')
            with rasterio.open(copol_tif) as copol:
                copol_value,row2,col2 = util_fcns.value_at(copol,lon1,lat1)
                copol_values.append(copol_value)
            with rasterio.open(crosspol_tif) as crosspol:
                crosspol_value,row3,col3 = util_fcns.value_at(
                    crosspol,lon1,lat1)
                crosspol_values.append(crosspol_value)

            # Make mosaic of Dist-S1 pngs
            quad_s1_png = target_name + '/quad_s1.png'
            img1 = cv2.imread(copol_png)
            img2 = cv2.imread(crosspol_png)
            img3 = cv2.imread(metric_pngs[-1])
            img4 = cv2.imread(status_pngs[-1])
            combined = plot_fcns.make_img_grid([img1,img2,img3,img4],2,2)
            cv2.imwrite(quad_s1_png,combined)

    # Dist-HLS pngs and data extraction
    dist_hls_status_pngs = []
    dt_dist_hls = []
    status_hls = []
    for df_prod_row in df_dist_hls_prod.itertuples(index=True):
        #if df_prod_row.Index == 0:
        print(df_prod_row.post_date)
        dt_dist_hls.append(pd.to_datetime(df_prod_row.post_date))
        filename = df_prod_row.prod_name / (
            df_prod_row.prod_name.name + '_' + prod_names_dist_hls[0])
        with rasterio.open(filename) as src:
            value,row1,col1 = util_fcns.value_at(src,lon1,lat1)
            status_hls.append(value)

        ofile = df_prod_row.prod_name / 'veg_dist_status.tif'
        ofile_png = ofile.with_suffix(".png")
        if redo_dist_hls_pngs:
            # Reproject geotif into same destination geometry
            ifile = df_prod_row.prod_name / (
                df_prod_row.prod_name.name + '_' + prod_names_s1[0])
            dst_width,dst_height = utils_geotif.reproject_to_geom(
                ifile,geo_crs,bbox,dst_res,ofile)
            # Make pngs
            titlestr = 'Dist-HLS ' + df_prod_row.post_date
            status_bounds2 = utils_geotif.geotif_to_png_map2(
                ofile,
                ofile_png,
                geo_crs,
                titlestr,
                DIST_STATUS_CMAP,
                0,6,lat1,lon1,circ_size,None)
        dist_hls_status_pngs.append(ofile_png)

    # HLS pngs and data extraction
    hls_pngs = []
    dt_hls = []
    values_hls = []
    for df_prod_row in df_hls_prod.itertuples(index=True):
        #if df_prod_row.Index == 0:
        print(df_prod_row.post_date)
        dt_hls.append(pd.to_datetime(df_prod_row.post_date,
            format="%Y%jT%H%M%S"))
        filename = df_prod_row.prod_name / (
            df_prod_row.prod_name.name + '.' + prod_names_hls[0])
        with rasterio.open(filename) as src:
            value,row1,col1 = util_fcns.value_at(src,lon1,lat1)
            values_hls.append(value)

        ofile = df_prod_row.prod_name / 'b03.tif'
        ofile_png = ofile.with_suffix(".png")
        if redo_hls_pngs:
            # Reproject geotif into same destination geometry
            ifile = df_prod_row.prod_name / (
                df_prod_row.prod_name.name + '.' + prod_names_hls[0])
            dst_width,dst_height = utils_geotif.reproject_to_geom(
                ifile,geo_crs,bbox,dst_res,ofile)
            # Substitute no data values to bottom of scale
            with rasterio.open(ofile) as src:
                data = src.read(1)
                udata = np.unique(data)
                mask = (data == udata[0])
                data[mask] = udata[1]-1
            util_fcns.modify_tifdata(ofile,data)
            datamin = data.min()
            datamax = data.max()
            if 0.5*datamax > datamin:
                datamax = 0.5*datamax

            # Make pngs
            titlestr = df_prod_row.post_date
            titlestr = 'HLS ' + df_prod_row.post_date
            status_bounds2 = utils_geotif.geotif_to_png_map2(
                ofile,
                ofile_png,
                geo_crs,
                titlestr,
                'gray',
                datamin,datamax,lat1,lon1,circ_size,None)
        hls_pngs.append(ofile_png)

    # Plot values at selected site
    print('plotting values')
    fpath = s1_pathbase / ('metric.png')
    titlestr = mgrs_tile_id
    plot_fcns.plot_val(fpath,dt_s1,copol_values,titlestr,"metric")
    fpath = s1_pathbase / ('copol.png')
    plot_fcns.plot_val_rtc(fpath,dt_s1,copol_values,titlestr)
    fpath = s1_pathbase / ('crosspol.png')
    plot_fcns.plot_val_rtc(fpath,dt_s1,crosspol_values,titlestr)
    fpath = hls_pathbase / ('hls.png')
    plot_fcns.plot_val_rtc(fpath,dt_hls,values_hls,titlestr)
    # mp4 at post-dates
    out_metric_mp4 = s1_pathbase / ('metric.mp4')
    out_metric_gif = s1_pathbase / ('metric.gif')
    out_copol_mp4 = s1_pathbase / ('copol.mp4')
    out_copol_gif = s1_pathbase / ('copol.gif')
    out_crosspol_mp4 = s1_pathbase / ('crosspol.mp4')
    out_crosspol_gif = s1_pathbase / ('crosspol.gif')
    out_hls_mp4 = hls_pathbase / ('hls.mp4')
    out_hls_gif = hls_pathbase / ('hls.gif')
    print(out_metric_mp4)
    plot_fcns.pngs_to_mp4(metric_pngs,out_metric_mp4,fps=mp4_rate)
    plot_fcns.pngs_to_gif(metric_pngs,out_metric_gif,fps=mp4_rate)
    print(out_copol_mp4)
    plot_fcns.pngs_to_mp4(copol_pngs,out_copol_mp4,fps=mp4_rate)
    plot_fcns.pngs_to_gif(copol_pngs,out_copol_gif,fps=mp4_rate)
    print(out_crosspol_mp4)
    plot_fcns.pngs_to_mp4(crosspol_pngs,out_crosspol_mp4,fps=mp4_rate)
    plot_fcns.pngs_to_gif(crosspol_pngs,out_crosspol_gif,fps=mp4_rate)
    print(out_hls_mp4)
    plot_fcns.pngs_to_mp4(hls_pngs,out_hls_mp4,fps=mp4_rate)
    plot_fcns.pngs_to_gif(hls_pngs,out_hls_gif,fps=mp4_rate)

    if redo_pair_pngs:
        # Create merged mp4 with s1 and hls pairs closest in time
        df_pairs = util_fcns.merge_post_dates(df_s1_prod,df_dist_hls_prod)
    
        print('Pair Pngs')
        pair_pngs = [] 
        for pair in df_pairs.itertuples():
            png_path1 = (Path(df_s1_prod.loc[pair[1]].prod_name)
                / 'dist_status.png')
            png_path2 = (Path(df_dist_hls_prod.loc[pair[2]].prod_name)
                / 'veg_dist_status.png')
            png_combined_path = (Path(df_s1_prod.loc[pair[1]].prod_name)
                / 'combined_veg.png')
            pair_pngs.append(png_combined_path)
            img1 = cv2.imread(png_path1)
            img2 = cv2.imread(png_path2)
            combined = cv2.vconcat([img1,img2])
            cv2.imwrite(png_combined_path,combined)

    if redo_quad_pngs:
        # Create merged pngs with Dist-S1,HLS status and post-date inputs
        print('Quad Pngs')
        df_quads = util_fcns.merge_post_dates_three(
            df_s1_prod,df_dist_hls_prod,df_hls_prod)
        quad_pngs = []
        for quad in df_quads.itertuples():
            png_path1 = (Path(df_s1_prod.loc[quad[1]].prod_name)
                / 'dist_status.png')
            png_path2 = (Path(df_dist_hls_prod.loc[quad[2]].prod_name)
                / 'veg_dist_status.png')
            png_path3 = (Path(df_s1_prod.loc[quad[1]].prod_name)
                / 'copol.png')
            png_path4 = (Path(df_hls_prod.loc[quad[3]].prod_name)
                / 'b03.png')
            png_combined_path = (Path(df_s1_prod.loc[quad[1]].prod_name)
                / 'combined_quad.png')
            quad_pngs.append(png_combined_path)
            img1 = cv2.imread(png_path1)
            img2 = cv2.imread(png_path2)
            img3 = cv2.imread(png_path3)
            img4 = cv2.imread(png_path4)
            row1 = cv2.hconcat([img1,img3])
            row2 = cv2.hconcat([img2,img4])
            combined = cv2.vconcat([row1,row2])
            cv2.imwrite(png_combined_path,combined)

            png_combined_path2 = (Path(df_s1_prod.loc[quad[1]].prod_name)
                / 'combined_quad2.png')
            png_path3a = (Path(df_s1_prod.loc[quad[1]].prod_name)
                / 'metric.png')
            img3a = cv2.imread(png_path3a)
            row1a = cv2.hconcat([img1,img3a])
            combined2 = cv2.vconcat([row1a,row2])
            cv2.imwrite(png_combined_path2,combined2)
    
    # mp4 of pairs
    out_pair_mp4 = target_name + '/' + target_name + '_pair.mp4'
    out_pair_gif = target_name + '/' + target_name + '_pair.gif'
    print(out_pair_mp4)
    plot_fcns.pngs_to_mp4(pair_pngs,out_pair_mp4,fps=mp4_rate)
    plot_fcns.pngs_to_gif(pair_pngs,out_pair_gif,fps=mp4_rate)

    # mp4 of quads
    out_quad_mp4 = target_name + '/' + target_name + '_quad.mp4'
    out_quad_gif = target_name + '/' + target_name + '_quad.gif'
    out_quad2_mp4 = target_name + '/' + target_name + '_quad2.mp4'
    out_quad2_gif = target_name + '/' + target_name + '_quad2.gif'
    print(out_quad_mp4)
    plot_fcns.pngs_to_mp4(quad_pngs,out_quad_mp4,fps=mp4_rate)
    plot_fcns.pngs_to_gif(quad_pngs,out_quad_gif,fps=mp4_rate)
    plot_fcns.pngs_to_mp4(quad_pngs,out_quad2_mp4,fps=mp4_rate)
    plot_fcns.pngs_to_gif(quad_pngs,out_quad2_gif,fps=mp4_rate)

    print("Building powerpoint slides")
    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    left = Inches(0.5)
    top = Inches(1)
    slide = prs.slides.add_slide(blank_slide_layout)
    textbox = slide.shapes.add_textbox(left,Inches(0.5),Inches(8),Inches(0.5))
    text_frame = textbox.text_frame
    text_frame.text = mgrs_tile_id
    pic = slide.shapes.add_picture(str(prod_png),left,top)
    rtc_pngs = [copol_png,crosspol_png]
    plot_fcns.add_image_grid_slides(prs,rtc_pngs,grid=(1,2))
    #slide = prs.slides.add_slide(blank_slide_layout)
    #pic = slide.shapes.add_picture(str(copol_png),left,top)
    #slide = prs.slides.add_slide(blank_slide_layout)
    #pic = slide.shapes.add_picture(str(crosspol_png),left,top)
    #slide = prs.slides.add_slide(blank_slide_layout)
    #pic = slide.shapes.add_picture(str(mask_png),left,top)
    prs.save(target_name + '_cmp_s1.pptx')

    raise Exception("End of main")

if __name__ == '__main__':
  try:
    main()
  except Exception as e:
    exc_type,exc_value,exc_traceback = sys.exc_info()
    tb = exc_traceback
    stack = inspect.trace()
    lcls = locals()
    dtb = debug_defs.to_doubly(tb)
    dtb1,sl,fname,lineno = debug_defs.locate_err_user(dtb,['opera','marimo'])
    if str(exc_value) == 'End of main':
        print(f"main ended normally at line: {lineno} in {fname}")
    else:
        print(f"Error at user line: {lineno} in {fname}")
        print(exc_value)
    # Put local variables from error scope into global scope
    locals_dict = dict(sl)
    for k,v in locals_dict.items():
        globals()[k] = v

